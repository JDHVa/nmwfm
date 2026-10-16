import asyncio
import shutil
import subprocess
import sys
import types
from datetime import date
from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from backend import config, services
from backend.core import diagramas, documents, ocr, producto

from .asgi_client import MiniClient
from .test_insumos_ocr import CtxRegistro, agregar_imagen
from .test_services import TID

MINDMAP = "mindmap\n  root((Tema))\n    Rama A\n    Rama B"


def falso_pdf(docx_path, pdf_path):
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_path.write_bytes(b"%PDF")
    return pdf_path


def texto_docx(ruta) -> str:
    return "".join(t.text for t in docx.Document(str(ruta)).element.iter(qn("w:t")))


def construir(seeded, monkeypatch, portada=None):
    monkeypatch.setattr(documents, "export_pdf", falso_pdf)
    seeded.storage.set_settings({"nombre": "Ana", "matricula": "1"})
    tarea = seeded.storage.get_tarea(TID)
    materia = seeded.storage.get_materia("100")
    return services._construir_documentos(
        seeded, tarea, materia, seeded.storage.get_settings(), "## A\n\nTexto", portada
    )


def test_portada_override_en_la_generacion(seeded, monkeypatch):
    gen = construir(seeded, monkeypatch, {"profesor": "DRA. NUEVA", "titulo": "Mi título propio"})
    contenido = texto_docx(gen["docx_path"])
    assert "DRA. NUEVA" in contenido and "PROFE UNO" not in contenido
    primera = docx.Document(gen["docx_path"]).paragraphs
    assert any(p.text == "Mi título propio" for p in primera)
    publico = seeded.generacion_public(gen)
    assert publico["portada"] == {"profesor": "DRA. NUEVA", "titulo": "Mi título propio"}


def test_sin_override_usa_materia_y_titulo(seeded, monkeypatch):
    gen = construir(seeded, monkeypatch)
    assert "PROFE UNO" in texto_docx(gen["docx_path"])
    assert seeded.generacion_public(gen)["portada"] is None


def test_restaurar_conserva_la_portada(seeded, monkeypatch):
    gen = construir(seeded, monkeypatch, {"profesor": "X"})
    seeded.storage.set_generacion(TID, "otra", gen["docx_path"], gen["pdf_path"], portada=None)
    antigua = seeded.storage.listar_generaciones(TID)[-1]
    restaurada = services.restaurar_generacion(seeded, antigua["id"])
    assert restaurada["portada"] == {"profesor": "X"}


def test_limpia_portada_vacia():
    assert services._portada_limpia({"profesor": "  ", "titulo": None}) is None
    assert services._portada_limpia({"profesor": " A ", "extra": "z"}) == {"profesor": "A"}


def test_api_generar_pasa_la_portada(app, seeded, monkeypatch):
    recibido = {}

    async def falso(state, ctx, tarea_id, extra, portada=None):
        recibido["portada"] = portada
        return {}

    monkeypatch.setattr(services, "run_generar", falso)
    client = MiniClient(app, headers={"X-Token": "dev"})

    async def escenario():
        r = await client.arequest(
            "POST",
            f"/tareas/{TID}/generar",
            json={"instrucciones_extra": "x", "portada": {"profesor": "P", "titulo": "T"}},
        )
        await asyncio.sleep(0.1)
        return r

    assert asyncio.run(escenario()).status_code == 200
    assert recibido["portada"] == {"profesor": "P", "titulo": "T"}


def test_portada_preview_con_overrides(app, seeded, monkeypatch):
    seeded.storage.set_settings({"nombre": "Ana", "matricula": "1"})
    capturado = {}

    def pdf(docx_path, pdf_path):
        capturado["texto"] = texto_docx(docx_path)
        return falso_pdf(docx_path, pdf_path)

    monkeypatch.setattr(documents, "export_pdf", pdf)
    client = MiniClient(app, headers={"X-Token": "dev"})
    r = client.get(f"/tareas/{TID}/portada-preview?profesor=DR.%20OVERRIDE")
    assert r.status_code == 200
    assert "DR. OVERRIDE" in capturado["texto"] and "PROFE UNO" not in capturado["texto"]
    client.get(f"/tareas/{TID}/portada-preview")
    assert "PROFE UNO" in capturado["texto"]


def test_patch_ocr_texto_marca_editado(app, seeded, tmp_path):
    fila = agregar_imagen(seeded, tmp_path)
    client = MiniClient(app, headers={"X-Token": "dev"})
    r = client.patch(f"/insumos/{fila['id']}", json={"ocr_texto": "texto corregido"})
    assert r.status_code == 200
    assert r.json()["ocr_texto"] == "texto corregido" and r.json()["ocr_editado"] is True


def test_ocr_editado_sobrevive_a_usar_recorte_pero_no_a_rotar(seeded, tmp_path):
    fila = agregar_imagen(seeded, tmp_path)
    st = seeded.storage
    st.update_insumo(fila["id"], ocr_texto="editado")
    st.update_insumo(fila["id"], usar_recorte=False)
    assert st.get_insumo(fila["id"])["ocr_texto"] == "editado"
    st.update_insumo(fila["id"], rotacion=90)
    actual = st.get_insumo(fila["id"])
    assert actual["ocr_texto"] is None and actual["ocr_editado"] == 0


def test_rotar_y_editar_a_la_vez_gana_la_edicion(seeded, tmp_path):
    fila = agregar_imagen(seeded, tmp_path)
    seeded.storage.update_insumo(fila["id"], rotacion=180, ocr_texto="manual")
    actual = seeded.storage.get_insumo(fila["id"])
    assert actual["ocr_texto"] == "manual" and actual["ocr_editado"] == 1


def test_patch_ocr_en_texto_da_400(app, seeded):
    fila = seeded.storage.add_insumo(TID, "texto", texto="x")
    client = MiniClient(app, headers={"X-Token": "dev"})
    assert client.patch(f"/insumos/{fila['id']}", json={"ocr_texto": "a"}).status_code == 400


def test_el_prompt_usa_el_texto_editado_aun_sin_ocr(seeded, tmp_path, monkeypatch):
    fila = agregar_imagen(seeded, tmp_path)
    seeded.storage.update_insumo(fila["id"], ocr_texto="Ley de Ohm corregida")
    monkeypatch.setattr(ocr, "disponible", lambda: False)
    texto = asyncio.run(services._texto_ocr(seeded, CtxRegistro(), TID))
    assert texto == "Ley de Ohm corregida"


class _Resultado:
    returncode = 0
    stdout = b"\x89PNGfalso"


def test_dot_por_ruta_absoluta_si_graphviz_falla(monkeypatch):
    class Roto:
        def __init__(self, *a, **kw):
            pass

        def pipe(self, **kw):
            raise RuntimeError("failed to execute dot")

    llamadas = []

    def run(cmd, **kw):
        llamadas.append(cmd[0])
        return _Resultado()

    monkeypatch.setitem(sys.modules, "graphviz", types.SimpleNamespace(Source=Roto))
    monkeypatch.setattr(shutil, "which", lambda nombre: None)
    monkeypatch.setattr(diagramas, "RUTAS_DOT", (str(Path(__file__)),))
    monkeypatch.setattr(subprocess, "run", run)
    assert diagramas.renderizar_png("digraph{a->b}").startswith(b"\x89PNG")
    assert llamadas == [str(Path(__file__))]


def test_sin_dot_en_ningun_lado_devuelve_none(monkeypatch):
    class Roto:
        def __init__(self, *a, **kw):
            pass

        def pipe(self, **kw):
            raise RuntimeError("failed to execute dot")

    monkeypatch.setitem(sys.modules, "graphviz", types.SimpleNamespace(Source=Roto))
    monkeypatch.setattr(shutil, "which", lambda nombre: None)
    monkeypatch.setattr(diagramas, "RUTAS_DOT", ("Z:\\no\\existe\\dot.exe",))
    assert diagramas.renderizar_png("digraph{a->b}") is None


@pytest.mark.skipif(diagramas.buscar_dot() is None, reason="Graphviz no está instalado")
def test_mindmap_real_con_dot():
    grafo = diagramas.parsear_mermaid(MINDMAP)
    png = diagramas.renderizar_png(diagramas.a_dot(grafo))
    assert png is not None and png.startswith(b"\x89PNG") and len(png) > 500


def test_prompts_se_resuelven_desde_resource_dir_empaquetado(tmp_path, monkeypatch):
    carpeta = tmp_path / "backend" / "core" / "prompts"
    carpeta.mkdir(parents=True)
    (carpeta / "generico.md").write_text("Tipo de producto: empaquetado.", encoding="utf-8")
    monkeypatch.setattr(config, "RESOURCE_DIR", tmp_path)
    assert producto.plantilla_producto("ensayo") == "Tipo de producto: empaquetado."
    assert config.directorio_prompts() == carpeta


def test_en_desarrollo_los_prompts_existen_bajo_resource_dir():
    assert (config.directorio_prompts() / "generico.md").exists()
    assert date.today()
