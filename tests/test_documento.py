from datetime import date
from pathlib import Path

import docx
from docx.oxml.ns import qn

from backend import config, services
from backend.core import documents

from .asgi_client import MiniClient


def secciones(n: int) -> str:
    return "\n\n".join(f"## Sección {i}\n\nTexto {i}." for i in range(1, n + 1))


def xml_texto(path) -> str:
    return docx.Document(str(path)).element.xml


def test_indice_solo_con_mas_de_cuatro_secciones(tmp_path):
    con = documents.markdown_to_docx(secciones(5), tmp_path / "a.docx")
    sin = documents.markdown_to_docx(secciones(4), tmp_path / "b.docx")
    assert "TOC" in xml_texto(con) and "Índice" in xml_texto(con)
    assert "TOC" not in xml_texto(sin) and "Índice" not in xml_texto(sin)
    textos = [p.text for p in docx.Document(str(con)).paragraphs]
    assert textos.count("Sección 3") >= 2


def test_referencias_con_sangria_francesa(tmp_path):
    md = (
        "## Desarrollo\n\nTexto normal.\n\n## Referencias\n\n"
        "Autor, A. (2020). *Libro uno*. Editorial.\n\n"
        "- Autor, B. (2021). Libro dos. Editorial.\n"
        "- Autor, C. (2022). Libro tres. Editorial.\n"
    )
    d = docx.Document(str(documents.markdown_to_docx(md, tmp_path / "r.docx")))
    refs = [p for p in d.paragraphs if p.text.startswith("Autor")]
    assert len(refs) == 3
    for p in refs:
        assert abs(p.paragraph_format.left_indent.cm - 1.27) < 0.01
        assert abs(p.paragraph_format.first_line_indent.cm + 1.27) < 0.01
        assert p.style.name != "List Bullet"
    normal = next(p for p in d.paragraphs if p.text == "Texto normal.")
    assert normal.paragraph_format.first_line_indent is None


def test_referencias_termina_con_la_siguiente_seccion(tmp_path):
    md = "## Referencias\n\nAutor, A. (2020). Libro.\n\n## Anexo\n\nTexto de anexo."
    d = docx.Document(str(documents.markdown_to_docx(md, tmp_path / "r.docx")))
    anexo = next(p for p in d.paragraphs if p.text == "Texto de anexo.")
    assert anexo.paragraph_format.first_line_indent is None


def test_encabezado_primera_linea_del_cuerpo(tmp_path):
    out = documents.markdown_to_docx("## A\n\nx", tmp_path / "e.docx", encabezado="Lógica - Ficha")
    assert docx.Document(str(out)).paragraphs[0].text == "Lógica - Ficha"


def test_numeracion_sin_numero_en_la_portada(tmp_path):
    mapping = documents.marcadores("ANA", "1", "PROFE", date(2026, 10, 14))
    out = documents.build_final_docx(
        config.PORTADA_TEMPLATE,
        "## A\n\nTexto",
        tmp_path / "f.docx",
        mapping,
        tmp_path / "tmp",
        encabezado="Mat - Tarea",
    )
    d = docx.Document(str(out))
    assert len(d.sections) >= 2
    cuerpo = d.sections[-1]
    assert not cuerpo.footer.is_linked_to_previous
    assert "PAGE" in cuerpo.footer._element.xml
    assert "PAGE" not in d.sections[0].footer._element.xml
    pg = cuerpo._sectPr.find(qn("w:pgNumType"))
    assert pg is not None and pg.get(qn("w:start")) == "1"


def gen(state, tarea_id, texto, docx_path=None, historial=True):
    base = Path(state.paths.generaciones)
    return state.storage.set_generacion(
        tarea_id,
        texto,
        str(docx_path or base / f"{texto}.docx"),
        str(base / f"{texto}.pdf"),
        [],
        historial=historial,
    )


def test_historial_de_generaciones_y_restaurar(seeded):
    st = seeded.storage
    gen(seeded, "100-1.1", "v1")
    gen(seeded, "100-1.1", "v2")
    lista = services.listar_generaciones(seeded, "100-1.1")
    assert [g["markdown"] for g in lista] == ["v2", "v1"]
    assert [g["actual"] for g in lista] == [True, False]
    assert lista[0]["id"] > lista[1]["id"]
    antigua = services.restaurar_generacion(seeded, lista[1]["id"])
    assert antigua["markdown"] == "v1" and antigua["id"] == lista[1]["id"]
    assert st.get_generacion("100-1.1")["markdown"] == "v1"
    assert [g["actual"] for g in services.listar_generaciones(seeded, "100-1.1")] == [False, True]


def test_reconstruir_no_agrega_al_historial(seeded):
    gen(seeded, "100-1.1", "v1")
    gen(seeded, "100-1.1", "v1", historial=False)
    assert len(services.listar_generaciones(seeded, "100-1.1")) == 1


def test_restaurar_inexistente_da_404(seeded):
    try:
        services.restaurar_generacion(seeded, 999)
    except services.ApiError as exc:
        assert exc.status == 404
    else:
        raise AssertionError("debía fallar")


def test_api_generaciones_y_restaurar(app, seeded):
    client = MiniClient(app, headers={"X-Token": "dev"})
    gen(seeded, "100-1.1", "v1")
    gen(seeded, "100-1.1", "v2")
    lista = client.get("/tareas/100-1.1/generaciones").json()
    assert len(lista) == 2 and lista[0]["actual"] is True
    r = client.post(f"/generaciones/{lista[1]['id']}/restaurar")
    assert r.status_code == 200 and r.json()["markdown"] == "v1"
    assert client.post("/generaciones/12345/restaurar").status_code == 404
    assert client.get("/tareas/no-existe/generaciones").status_code == 404


def test_api_portada_preview_con_datos_reales(app, seeded, monkeypatch):
    seeded.storage.set_settings({"nombre": "Ana Prueba", "matricula": "1999"})
    capturado = {}

    def falso_pdf(docx_path, pdf_path):
        capturado["texto"] = "".join(
            t.text for t in docx.Document(str(docx_path)).element.iter(qn("w:t"))
        )
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        pdf_path.write_bytes(b"%PDF")
        return pdf_path

    monkeypatch.setattr(documents, "export_pdf", falso_pdf)
    client = MiniClient(app, headers={"X-Token": "dev"})
    r = client.get("/tareas/100-1.1/portada-preview")
    assert r.status_code == 200 and r.json()["pdf_url"].endswith("portada.pdf")
    assert "Ana Prueba" in capturado["texto"] and "PROFE UNO" in capturado["texto"]


def test_api_portada_preview_sin_word(app, seeded, monkeypatch):
    def falla(docx_path, pdf_path):
        raise RuntimeError("Word no disponible")

    monkeypatch.setattr(documents, "export_pdf", falla)
    client = MiniClient(app, headers={"X-Token": "dev"})
    r = client.get("/tareas/100-1.1/portada-preview")
    assert r.status_code == 502 and r.json()["error"]["code"] == "sin_word"
