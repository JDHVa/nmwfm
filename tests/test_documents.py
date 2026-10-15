from datetime import date

import docx
from docx.oxml.ns import qn

from backend import config
from backend.core import documents


def all_text(path) -> str:
    d = docx.Document(str(path))
    return "\n".join(
        "".join(t.text for t in p.iter(qn("w:t"))) for p in d.element.body.iter(qn("w:p"))
    )


def test_fecha_es():
    assert documents.fecha_es(date(2026, 10, 14)) == "14 de octubre del 2026"
    assert documents.fecha_es(date(2026, 1, 3)) == "3 de enero del 2026"


def test_fill_cover_replaces_split_markers(tmp_path):
    mapping = documents.marcadores("ANA LOPEZ", "1999999", "PROFE UNO", date(2026, 10, 14))
    out = documents.fill_cover(config.PORTADA_TEMPLATE, tmp_path / "p.docx", mapping)
    text = all_text(out)
    for marker in mapping:
        assert marker not in text
    assert "ANA LOPEZ" in text and "1999999" in text and "PROFE UNO" in text
    assert "A 14 DE octubre DEL 2026" in text


def test_markdown_to_docx_styles(tmp_path):
    md = (
        "## Titulo\n\nTexto **fuerte** y *suave*.\n\n- a\n  - b\n\n1. uno\n2. dos\n\n"
        "| X | Y |\n|---|---|\n| 1 | 2 |\n"
    )
    out = documents.markdown_to_docx(md, tmp_path / "b.docx")
    d = docx.Document(str(out))
    assert d.sections[0].left_margin.cm == 2.54
    normal = d.styles["Normal"]
    assert normal.font.name == "Arial" and normal.font.size.pt == 12
    assert normal.paragraph_format.line_spacing == 1.5
    assert len(d.tables) == 1
    strong = [r for p in d.paragraphs for r in p.runs if r.text == "fuerte"]
    assert strong and strong[0].bold
    assert any(p.style.name == "Heading 2" for p in d.paragraphs)
    assert any(p.style.name.startswith("List Bullet") for p in d.paragraphs)


def test_build_final_docx(tmp_path):
    mapping = documents.marcadores("ANA LOPEZ", "1999999", "PROFE UNO", date(2026, 10, 14))
    out = documents.build_final_docx(
        config.PORTADA_TEMPLATE, "## Hola\n\nCuerpo", tmp_path / "f.docx", mapping, tmp_path / "w"
    )
    d = docx.Document(str(out))
    text = all_text(out)
    assert text.index("ANA LOPEZ") < text.index("Hola") < text.index("Cuerpo")
    assert len(d.sections) == 2
    assert d.sections[-1].left_margin.cm == 2.54
