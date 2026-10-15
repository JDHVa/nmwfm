import io

import docx
from PIL import Image

from backend.core import diagramas, documents
from backend.core.prompt_builder import PromptData, build_prompt

MINDMAP = """```mermaid
mindmap
  root((Fotosíntesis))
    Fase luminosa
      Tilacoides
      ATP
    Fase oscura
      Ciclo de Calvin
```"""

FLUJO = """```mermaid
graph LR
  A[Idea general] --> B[Categoría 1]
  A --> C["Categoría 2"]
  B -->|incluye| D(Detalle)
  C -- depende --> D
```"""

TIMELINE = """```mermaid
timeline
    title Independencia
    1810 : Grito de Dolores
    1821 : Consumación : Entrada del Ejército Trigarante
```"""


def codigo(bloque: str) -> str:
    return diagramas.extraer_bloques(bloque)[0]


def png_falso() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(buf, "PNG")
    return buf.getvalue()


def test_mindmap_arbol_y_lista():
    g = diagramas.parsear_mermaid(codigo(MINDMAP))
    assert g.tipo == "mindmap" and len(g.nodos) == 6
    assert ("Fotosíntesis" in g.nodos.values()) and len(g.aristas) == 5
    lista = diagramas.a_markdown(g)
    assert "- Fotosíntesis\n  - Fase luminosa\n    - Tilacoides" in lista


def test_flujo_nodos_etiquetas_y_enlaces_con_texto():
    g = diagramas.parsear_mermaid(codigo(FLUJO))
    assert g.direccion == "LR"
    assert g.nodos["A"] == "Idea general" and g.nodos["C"] == "Categoría 2"
    assert g.nodos["D"] == "Detalle"
    assert ("B", "D", "incluye") in g.aristas and ("C", "D", "depende") in g.aristas
    tabla = diagramas.a_markdown(g)
    assert "| Idea general | → | Categoría 1 |" in tabla


def test_timeline_en_cadena():
    g = diagramas.parsear_mermaid(codigo(TIMELINE))
    assert len(g.nodos) == 2 and len(g.aristas) == 1
    assert "Entrada del Ejército Trigarante" in " ".join(g.nodos.values())


def test_dot_escapa_y_contiene_aristas():
    g = diagramas.parsear_mermaid('graph TD\n  A["Dice \\"hola\\""] --> B')
    dot = diagramas.a_dot(g)
    assert dot.startswith("digraph G {") and '"A" -> "B"' in dot and "rankdir=TB" in dot


def test_codigo_invalido_devuelve_none():
    assert diagramas.parsear_mermaid("pie title X\n  a: 1") is None
    assert diagramas.parsear_mermaid("") is None


def test_sin_renderizador_deja_lista_y_avisa(tmp_path, monkeypatch):
    monkeypatch.setattr(diagramas, "renderizar_png", lambda dot: None)
    avisos = []
    md, imagenes = diagramas.procesar_diagramas("Texto\n\n" + MINDMAP, tmp_path, avisos)
    assert imagenes == {} and "- Fotosíntesis" in md and "```" not in md
    assert len(avisos) == 1 and "Graphviz" in avisos[0]


def test_con_renderizador_inserta_imagen_en_el_word(tmp_path, monkeypatch):
    monkeypatch.setattr(diagramas, "renderizar_png", lambda dot: png_falso())
    avisos = []
    salida = documents.markdown_to_docx(
        "## Mapa\n\nIntro\n\n" + MINDMAP + "\n\nCierre", tmp_path / "b.docx", tmp_path, avisos
    )
    d = docx.Document(str(salida))
    assert len(d.inline_shapes) == 1 and avisos == []
    textos = "\n".join(p.text for p in d.paragraphs)
    assert "[[DIAGRAMA" not in textos and "Intro" in textos and "Cierre" in textos


def test_word_sin_renderizador_usa_lista_y_registra_aviso(tmp_path, monkeypatch):
    monkeypatch.setattr(diagramas, "renderizar_png", lambda dot: None)
    avisos = []
    salida = documents.markdown_to_docx(FLUJO, tmp_path / "b.docx", tmp_path, avisos)
    d = docx.Document(str(salida))
    assert len(d.inline_shapes) == 0 and len(d.tables) == 1 and len(avisos) == 1


def test_renderizado_real_si_hay_graphviz():
    import shutil

    if shutil.which("dot") is None:
        return
    png = diagramas.renderizar_png(diagramas.a_dot(diagramas.parsear_mermaid(codigo(FLUJO))))
    assert png and png.startswith(b"\x89PNG")


def test_prompt_pide_mermaid_solo_en_tipos_de_diagrama():
    con = build_prompt(PromptData(materia="M", titulo="Dimensión 3 - Mapa mental"))
    sin = build_prompt(PromptData(materia="M", titulo="Dimensión 2 - Ficha Técnica"))
    assert "```mermaid" in con and "mindmap" in con
    assert "mermaid" not in sin
    linea = build_prompt(PromptData(materia="M", titulo="Línea del tiempo"))
    assert "timeline" in linea


def test_avisos_se_guardan_y_se_exponen(seeded):
    D = str(seeded.paths.generaciones / "a.bin")
    seeded.storage.set_generacion("100-1.1", "md", D, D, ["Diagrama 1: sin Graphviz"])
    pub = seeded.generacion_public(seeded.storage.get_generacion("100-1.1"))
    assert pub["avisos"] == ["Diagrama 1: sin Graphviz"]
    seeded.storage.set_generacion("100-1.1", "md", D, D)
    assert seeded.generacion_public(seeded.storage.get_generacion("100-1.1"))["avisos"] == []
