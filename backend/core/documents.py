import unicodedata
from datetime import date
from pathlib import Path

import mistune
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer

from . import diagramas

MESES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)
FONT = "Arial"
BODY_PT = 12
LINE_SPACING = 1.5
MARGIN_CM = 2.54
HEADING_PT = {1: 16, 2: 14, 3: 13}
SECCIONES_PARA_INDICE = 4
HUECO_APA_CM = 1.27
TITULOS_REFERENCIAS = (
    "referencias",
    "referencias bibliograficas",
    "referencias apa",
    "bibliografia",
    "fuentes consultadas",
)


def _plano(texto: str) -> str:
    nfd = unicodedata.normalize("NFD", texto)
    return "".join(c for c in nfd if unicodedata.category(c) != "Mn").lower().strip(" :.")


def es_titulo_referencias(texto: str) -> bool:
    return _plano(texto) in TITULOS_REFERENCIAS


def _texto_tokens(tokens) -> str:
    partes = []
    for tok in tokens:
        if tok.get("children"):
            partes.append(_texto_tokens(tok["children"]))
        else:
            partes.append(tok.get("raw", ""))
    return "".join(partes)


def _campo(run, tipo: str) -> None:
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), tipo)
    run._r.append(el)


def _instruccion(run, texto: str) -> None:
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = texto
    run._r.append(el)


def fecha_es(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} del {d.year}"


def marcadores(nombre: str, matricula: str, profesor: str, hoy: date) -> dict[str, str]:
    return {
        "(NOMBRE DE LA PERSONA)": nombre,
        "(MATRICULA DE LA PERSONA)": matricula,
        "(NOMBRE DEL PROFE)": profesor,
        "(DIA XX)": str(hoy.day),
        "(MES)": MESES[hoy.month - 1].upper(),
        "(AÑO)": str(hoy.year),
    }


def _replace_in_paragraph(paragraph: Paragraph, mapping: dict[str, str]) -> bool:
    runs = paragraph.runs
    texts = [r.text for r in runs]
    if not any(key in "".join(texts) for key in mapping):
        return False
    for key, value in mapping.items():
        while (idx := "".join(texts).find(key)) != -1:
            end = idx + len(key)
            pos, first, last = 0, None, None
            for i, text in enumerate(texts):
                if first is None and idx < pos + len(text):
                    first, first_off = i, idx - pos
                if first is not None and end <= pos + len(text):
                    last, last_off = i, end - pos
                    break
                pos += len(text)
            if first == last:
                t = texts[first]
                texts[first] = t[:first_off] + value + t[last_off:]
            else:
                texts[first] = texts[first][:first_off] + value
                for j in range(first + 1, last):
                    texts[j] = ""
                texts[last] = texts[last][last_off:]
    for run, text in zip(runs, texts, strict=True):
        if run.text != text:
            run.text = text
    return True


def _all_paragraphs(doc: DocumentType):
    roots = [doc.element.body]
    for section in doc.sections:
        roots += [section.header._element, section.footer._element]
    for root in roots:
        for p in root.iter(qn("w:p")):
            yield Paragraph(p, None)


def fill_cover(template: Path, out: Path, mapping: dict[str, str]) -> Path:
    doc = Document(str(template))
    for paragraph in _all_paragraphs(doc):
        _replace_in_paragraph(paragraph, mapping)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def _set_rfonts(rpr_owner, name: str) -> None:
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def _style_font(style, size: int, bold: bool | None = None, name: str = FONT) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    _set_rfonts(style.element, name)


class _Builder:
    def __init__(self, imagenes: dict[str, Path] | None = None) -> None:
        self.imagenes = imagenes or {}
        self.en_referencias = False
        self.doc = Document()
        section = self.doc.sections[0]
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(MARGIN_CM))
        styles = self.doc.styles
        _style_font(styles["Normal"], BODY_PT)
        styles["Normal"].paragraph_format.line_spacing = LINE_SPACING
        styles["Normal"].paragraph_format.space_after = Pt(6)
        for level in range(1, 4):
            _style_font(styles[f"Heading {level}"], HEADING_PT[level], bold=True)
        for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
            _style_font(styles[name], BODY_PT)

    def run(self, paragraph, text, *, bold=False, italic=False, code=False, size=BODY_PT):
        run = paragraph.add_run(text)
        name = "Courier New" if code else FONT
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold or None
        run.font.italic = italic or None
        run.font.color.rgb = RGBColor(0, 0, 0)
        _set_rfonts(run._r, name)
        return run

    def inline(self, paragraph, tokens, *, bold=False, italic=False, size=BODY_PT):
        for tok in tokens:
            kind = tok["type"]
            if kind == "text":
                self.run(paragraph, tok["raw"], bold=bold, italic=italic, size=size)
            elif kind == "strong":
                self.inline(paragraph, tok["children"], bold=True, italic=italic, size=size)
            elif kind == "emphasis":
                self.inline(paragraph, tok["children"], bold=bold, italic=True, size=size)
            elif kind == "strikethrough":
                self.inline(paragraph, tok["children"], bold=bold, italic=italic, size=size)
            elif kind == "codespan":
                self.run(paragraph, tok["raw"], code=True, size=size)
            elif kind == "link":
                self.inline(paragraph, tok["children"], bold=bold, italic=italic, size=size)
            elif kind == "image":
                self.run(paragraph, tok.get("attrs", {}).get("alt", ""), size=size)
            elif kind == "softbreak":
                self.run(paragraph, " ", size=size)
            elif kind == "linebreak":
                paragraph.add_run().add_break()
            elif tok.get("children"):
                self.inline(paragraph, tok["children"], bold=bold, italic=italic, size=size)
            elif tok.get("raw"):
                self.run(paragraph, tok["raw"], bold=bold, italic=italic, size=size)

    def paragraph(self, tokens, style=None, indent_cm=0.0, spacing=LINE_SPACING):
        p = self.doc.add_paragraph(style=style)
        p.paragraph_format.line_spacing = spacing
        if indent_cm:
            p.paragraph_format.left_indent = Cm(indent_cm)
        self.inline(p, tokens)
        return p

    def blocks(self, tokens, depth=0):
        for tok in tokens:
            kind = tok["type"]
            if kind == "heading":
                level = tok["attrs"]["level"]
                self.en_referencias = es_titulo_referencias(_texto_tokens(tok["children"]))
                p = self.doc.add_paragraph(style=f"Heading {min(level, 3)}")
                p.paragraph_format.line_spacing = LINE_SPACING
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.keep_with_next = True
                self.inline(p, tok["children"], bold=True, size=HEADING_PT.get(level, BODY_PT))
            elif kind in ("paragraph", "block_text") and self._es_diagrama(tok):
                self.diagrama(tok)
            elif kind in ("paragraph", "block_text") and self.en_referencias:
                self.referencia(tok["children"])
            elif kind in ("paragraph", "block_text"):
                p = self.paragraph(tok["children"])
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif kind == "list":
                self.list(tok, depth)
            elif kind == "block_code":
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.5)
                self.run(p, tok["raw"].rstrip("\n"), code=True, size=10)
            elif kind == "block_quote":
                start = len(self.doc.paragraphs)
                self.blocks(tok["children"], depth)
                for p in self.doc.paragraphs[start:]:
                    p.paragraph_format.left_indent = Cm(1.0)
                    for r in p.runs:
                        r.font.italic = True
            elif kind == "table":
                self.table(tok)
            elif kind == "thematic_break":
                continue
            elif tok.get("children"):
                self.blocks(tok["children"], depth)

    def _es_diagrama(self, tok) -> bool:
        texto = "".join(c.get("raw", "") for c in tok["children"]).strip()
        return texto in self.imagenes

    def diagrama(self, tok) -> None:
        texto = "".join(c.get("raw", "") for c in tok["children"]).strip()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(self.imagenes[texto]), width=Cm(15))

    def referencia(self, hijos) -> None:
        p = self.paragraph(hijos)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(HUECO_APA_CM)
        p.paragraph_format.first_line_indent = Cm(-HUECO_APA_CM)

    def encabezado(self, texto: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.run(p, texto, bold=True, size=14)

    def indice(self, titulos: list[tuple[int, str]]) -> None:
        p = self.doc.add_paragraph()
        self.run(p, "Índice", bold=True, size=14)
        for k, (nivel, texto) in enumerate(titulos):
            p = self.doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.6 * (nivel - 1))
            if k == 0:
                _campo(p.add_run(), "begin")
                _instruccion(p.add_run(), ' TOC \\o "1-3" \\h \\z \\u ')
                _campo(p.add_run(), "separate")
            self.run(p, texto)
            if k == len(titulos) - 1:
                _campo(p.add_run(), "end")

    def list(self, tok, depth):
        if self.en_referencias:
            for item in tok["children"]:
                for child in item["children"]:
                    if child["type"] in ("block_text", "paragraph"):
                        self.referencia(child["children"])
            return
        ordered = tok["attrs"].get("ordered", False)
        start = tok["attrs"].get("start") or 1
        for n, item in enumerate(tok["children"], start=start):
            first = True
            for child in item["children"]:
                if child["type"] in ("block_text", "paragraph") and first:
                    first = False
                    if ordered:
                        p = self.paragraph(
                            [{"type": "text", "raw": f"{n}. "}, *child["children"]],
                            indent_cm=0.75 + 0.75 * depth,
                        )
                        p.paragraph_format.first_line_indent = Cm(-0.5)
                    else:
                        style = "List Bullet" if depth == 0 else f"List Bullet {min(depth + 1, 3)}"
                        self.paragraph(child["children"], style=style)
                elif child["type"] == "list":
                    self.list(child, depth + 1)
                else:
                    self.blocks([child], depth + 1)

    def table(self, tok):
        head = next((c for c in tok["children"] if c["type"] == "table_head"), None)
        body = next((c for c in tok["children"] if c["type"] == "table_body"), None)
        rows = []
        if head:
            rows.append((head["children"], True))
        if body:
            rows += [(r["children"], False) for r in body["children"]]
        if not rows:
            return
        cols = max(len(cells) for cells, _ in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for r, (cells, is_head) in enumerate(rows):
            for c, cell in enumerate(cells):
                target = table.cell(r, c)
                p = target.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                self.inline(p, cell["children"], bold=is_head, size=11)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)


def markdown_to_docx(
    markdown: str,
    out: Path,
    work_dir: Path | None = None,
    avisos: list[str] | None = None,
    encabezado: str | None = None,
) -> Path:
    imagenes: dict[str, Path] = {}
    if diagramas.extraer_bloques(markdown):
        markdown, imagenes = diagramas.procesar_diagramas(
            markdown, work_dir or out.parent / "diagramas", avisos
        )
    parser = mistune.create_markdown(renderer=None, plugins=["table", "strikethrough"])
    tokens = parser(markdown)
    builder = _Builder(imagenes)
    if encabezado:
        builder.encabezado(encabezado)
    titulos = [
        (min(t["attrs"]["level"], 3), _texto_tokens(t["children"]))
        for t in tokens
        if t["type"] == "heading"
    ]
    if sum(1 for nivel, _ in titulos if nivel == 2) > SECCIONES_PARA_INDICE:
        builder.indice(titulos)
    builder.blocks(tokens)
    out.parent.mkdir(parents=True, exist_ok=True)
    builder.doc.save(str(out))
    return out


def _numerar_paginas(section) -> None:
    section.footer.is_linked_to_previous = False
    parrafo = section.footer.paragraphs[0]
    for r in list(parrafo.runs):
        r._r.getparent().remove(r._r)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = parrafo.add_run()
    _campo(run, "begin")
    _instruccion(parrafo.add_run(), " PAGE ")
    _campo(parrafo.add_run(), "separate")
    parrafo.add_run("1")
    _campo(parrafo.add_run(), "end")
    sect = section._sectPr
    if sect.find(qn("w:pgNumType")) is None:
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        cols = sect.find(qn("w:cols"))
        if cols is not None:
            cols.addprevious(pg)
        else:
            sect.append(pg)


def combine(cover: Path, body: Path, out: Path) -> Path:
    cover_doc = Document(str(cover))
    section = cover_doc.add_section(WD_SECTION.NEW_PAGE)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))
    composer = Composer(cover_doc)
    composer.append(Document(str(body)))
    _numerar_paginas(composer.doc.sections[-1])
    out.parent.mkdir(parents=True, exist_ok=True)
    composer.save(str(out))
    return out


def build_final_docx(
    template: Path,
    markdown: str,
    out: Path,
    mapping: dict[str, str],
    work_dir: Path,
    avisos: list[str] | None = None,
    encabezado: str | None = None,
) -> Path:
    work_dir.mkdir(parents=True, exist_ok=True)
    cover = fill_cover(template, work_dir / "portada.docx", mapping)
    body = markdown_to_docx(markdown, work_dir / "cuerpo.docx", work_dir, avisos, encabezado)
    return combine(cover, body, out)


def export_pdf(docx_path: Path, pdf_path: Path) -> Path:
    from docx2pdf import convert

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    convert(str(docx_path), str(pdf_path))
    return pdf_path
